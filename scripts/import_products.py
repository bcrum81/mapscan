#!/usr/bin/env python3
"""Load the product spreadsheet into `products`.

Expected columns (header row, case-insensitive, order-independent):
    title, upc, sku, msrp, map, description

Validation happens here rather than at crawl time, because a transposed MAP
generates phantom violations forever and a bad UPC silently destroys matching.

    python -m scripts.import_products products.xlsx           # dry run
    python -m scripts.import_products products.xlsx --commit
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import db                                  # noqa: E402
from worker.match import normalize_upc, upc_check_digit_valid  # noqa: E402

CANONICAL = {
    "title": {"title", "product", "product title", "name", "product name"},
    "upc": {"upc", "gtin", "ean", "barcode", "upc code"},
    "sku": {"sku", "item", "item number", "part number", "mpn", "model"},
    "msrp": {"msrp", "retail", "retail price", "list", "list price"},
    "map": {"map", "map price", "min advertised price", "minimum advertised price"},
    "description": {"description", "desc", "details", "long description"},
    "wholesale": {"wholesale", "wholesale price", "cost", "dealer price",
                  "dealer cost", "your cost"},
}


def canonical_header(raw: str) -> str | None:
    key = (raw or "").strip().lower()
    for canon, aliases in CANONICAL.items():
        if key == canon or key in aliases:
            return canon
    return None


def promote_description_to_title(rows: list[dict]) -> bool:
    """Use Description as the title when the sheet has no title column.

    Real sheets often carry a single short product-name column called
    "Description" -- e.g. "Nova Glock X300U, Full w/ Soft Loops". Title is the
    field matching actually depends on (URL-slug pre-filtering, fuzzy title
    scoring, and fitment-token extraction all read it), so treating that column
    as a description and leaving title empty would skip every row.

    Only fires when there is no real title column, so a sheet with both keeps
    them distinct.
    """
    if not rows:
        return False
    has_title = any((r.get("title") or "").strip() for r in rows)
    has_desc = any((r.get("description") or "").strip() for r in rows)
    if has_title or not has_desc:
        return False
    for r in rows:
        r["title"] = r.get("description")
        # Not copied into description as well: the matcher compares our
        # description against the dealer page's description, and scoring a short
        # product name against a page of prose only adds noise.
        r["description"] = None
    return True


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        cols = [canonical_header(str(h)) if h is not None else None for h in header]
        out = []
        for r in rows:
            if all(v is None or str(v).strip() == "" for v in r):
                continue
            out.append({c: r[i] for i, c in enumerate(cols) if c})
        wb.close()
        return out

    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        mapping = {f: canonical_header(f) for f in (reader.fieldnames or [])}
        return [
            {mapping[k]: v for k, v in row.items() if mapping.get(k)}
            for row in reader
        ]


def money(value) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    text = str(value).replace("$", "").replace(",", "").strip()
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def validate(rows: list[dict]) -> tuple[list[dict], list[str]]:
    clean: list[dict] = []
    problems: list[str] = []
    seen_sku: dict[str, int] = {}

    for i, row in enumerate(rows, start=2):    # +2: header is row 1
        sku = str(row.get("sku") or "").strip()
        title = str(row.get("title") or "").strip()
        map_price = money(row.get("map"))
        msrp = money(row.get("msrp"))
        upc_raw = row.get("upc")
        upc = normalize_upc(str(upc_raw) if upc_raw is not None else None)

        if not sku:
            problems.append(f"row {i}: missing SKU — skipped")
            continue
        if not title:
            problems.append(f"row {i}: missing title — skipped")
            continue
        if map_price is None or map_price <= 0:
            problems.append(f"row {i} ({sku}): missing or invalid MAP — skipped")
            continue
        if sku in seen_sku:
            problems.append(f"row {i} ({sku}): duplicate SKU, first seen row {seen_sku[sku]} — skipped")
            continue
        seen_sku[sku] = i

        # These two are the ones that silently poison the system.
        if msrp is not None and map_price > msrp:
            problems.append(
                f"row {i} ({sku}): MAP {map_price} exceeds MSRP {msrp} — "
                "likely transposed columns; SKIPPED"
            )
            continue
        if upc_raw and not upc:
            problems.append(f"row {i} ({sku}): UPC {upc_raw!r} unparseable — importing without UPC")
        elif upc and not upc_check_digit_valid(upc):
            problems.append(
                f"row {i} ({sku}): UPC {upc} fails GS1 check digit — "
                "importing anyway, but verify it (bad UPCs break matching silently)"
            )

        clean.append({
            "sku": sku, "title": title, "upc": upc,
            "msrp": msrp, "map_price": map_price,
            "wholesale": money(row.get("wholesale")),
            "description": (str(row.get("description")).strip()
                            if row.get("description") else None),
        })

    # --- cross-row checks -------------------------------------------------
    # A UPC shared by two SKUs breaks UPC matching outright: a listing carrying
    # that UPC matches both products and there is no way to tell which. This is
    # a data problem only the owner can resolve, so it is surfaced loudly.
    upc_owners: dict[str, list[str]] = {}
    for r in clean:
        if r["upc"]:
            upc_owners.setdefault(r["upc"], []).append(r["sku"])
    for upc, skus in upc_owners.items():
        if len(skus) > 1:
            problems.append(
                f"DUPLICATE UPC {upc} shared by {', '.join(skus)} -- UPC matching "
                "cannot distinguish these; fix the sheet or clear one UPC"
            )

    equal_map = [r["sku"] for r in clean
                 if r["msrp"] is not None and r["map_price"] == r["msrp"]]
    if equal_map:
        shown = ", ".join(equal_map[:6]) + (" ..." if len(equal_map) > 6 else "")
        problems.append(
            f"{len(equal_map)} row(s) have MAP equal to MSRP ({shown}) -- "
            "verify that is intentional; any discount at all becomes a violation"
        )
    return clean, problems


async def upsert(rows: list[dict]) -> tuple[int, int]:
    await db.init_pool()
    inserted = updated = 0
    for r in rows:
        existed = await db.fetchval("SELECT 1 FROM products WHERE sku=$1", r["sku"])
        await db.execute(
            """
            INSERT INTO products (sku, title, upc, msrp, map_price, description,
                                  wholesale_price, active)
            VALUES ($1,$2,$3,$4,$5,$6,$7,true)
            ON CONFLICT (sku) DO UPDATE SET
                title = EXCLUDED.title,
                upc = EXCLUDED.upc,
                msrp = EXCLUDED.msrp,
                map_price = EXCLUDED.map_price,
                description = EXCLUDED.description,
                wholesale_price = EXCLUDED.wholesale_price,
                active = true,
                updated_at = now()
            """,
            r["sku"], r["title"], r["upc"], r["msrp"], r["map_price"],
            r["description"], r["wholesale"],
        )
        if existed:
            updated += 1
        else:
            inserted += 1
    await db.close_pool()
    return inserted, updated


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", type=Path, help=".xlsx or .csv")
    ap.add_argument("--commit", action="store_true",
                    help="write to the database (default is a dry run)")
    args = ap.parse_args()

    if not args.path.exists():
        print(f"no such file: {args.path}", file=sys.stderr)
        return 1

    raw = read_rows(args.path)
    promoted = promote_description_to_title(raw)
    clean, problems = validate(raw)

    print(f"read {len(raw)} data rows from {args.path.name}")
    if promoted:
        print("NOTE: no title column found; using 'Description' as the product "
              "title.\n      Matching depends on title, so this is required -- "
              "not a fallback.")
    print(f"valid: {len(clean)}   problems: {len(problems)}")
    if problems:
        print("\n--- problems ---")
        for p in problems:
            print(" ", p)

    missing_upc = sum(1 for r in clean if not r["upc"])
    if missing_upc:
        print(f"\nnote: {missing_upc}/{len(clean)} products have no UPC. UPC is the "
              "single strongest match signal — expect more manual review for those.")

    if not args.commit:
        print("\nDRY RUN — nothing written. Re-run with --commit to apply.")
        for r in clean[:5]:
            print(f"  sample: {r['sku']:14} | {r['title'][:44]:46} | MAP {r['map_price']}")
        return 0

    ins, upd = asyncio.run(upsert(clean))
    print(f"\ncommitted: {ins} inserted, {upd} updated")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
